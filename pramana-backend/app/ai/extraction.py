import csv
import json
import logging
import os
from typing import List, Optional, Dict, Any

import docx
import PyPDF2
from pydantic import BaseModel, Field

from app.config import settings

logger = logging.getLogger("pramana.ai.extraction")


class ExtractionError(Exception):
    """Base exception for document extraction failures."""
    pass


class UnsupportedFormatError(ExtractionError):
    """Raised when an unsupported or legacy file format is encountered."""
    pass


class EmptyExtractionError(ExtractionError):
    """Raised when extraction produces zero usable text (e.g. empty file or scanned image PDF)."""
    pass


class FileCorruptedError(ExtractionError):
    """Raised when a file cannot be parsed due to corruption."""
    pass


class FileSizeLimitExceededError(ExtractionError):
    """Raised when a file exceeds memory or upload size limits."""
    pass


class ExtractedSegment(BaseModel):
    page_number: Optional[int] = Field(default=None, description="1-indexed page number if available")
    section_name: Optional[str] = Field(default=None, description="Section or heading title")
    text: str = Field(description="Extracted textual content")


def extract_document(file_path: str, file_type: Optional[str] = None) -> List[ExtractedSegment]:
    """
    Extracts text from supported compliance artifacts, returning structured segments with page metadata.
    Supported formats: PDF, DOCX, TXT, JSON, CSV.
    Explicitly rejects legacy .doc files with clear instruction.
    Raises clear, typed exceptions for missing, empty, or corrupted files.
    """
    if not os.path.exists(file_path):
        raise ExtractionError(f"Evidence file not found on disk: {file_path}")

    file_size = os.path.getsize(file_path)
    if file_size == 0:
        raise EmptyExtractionError("Evidence file is completely empty (0 bytes).")

    if file_size > settings.max_upload_size_bytes:
        max_mb = settings.max_upload_size_bytes // (1024 * 1024)
        raise FileSizeLimitExceededError(
            f"File size ({file_size / (1024 * 1024):.1f} MB) exceeds maximum allowed limit ({max_mb} MB)."
        )

    # Determine file type
    if not file_type:
        ext = os.path.splitext(file_path)[1].lower().lstrip(".")
        file_type = ext
    else:
        file_type = file_type.lower().lstrip(".")

    # Reject legacy .doc explicitly
    if file_type == "doc":
        raise UnsupportedFormatError(
            "Legacy binary Word format (.doc) is not supported. "
            "Please convert the document to modern (.docx) or PDF before uploading."
        )

    segments: List[ExtractedSegment] = []

    # 1. PDF Documents
    if file_type == "pdf":
        try:
            with open(file_path, "rb") as f:
                reader = PyPDF2.PdfReader(f)
                num_pages = len(reader.pages)
                if num_pages == 0:
                    raise EmptyExtractionError("PDF contains zero pages.")

                for idx, page in enumerate(reader.pages, start=1):
                    try:
                        extracted = page.extract_text()
                    except Exception as e:
                        logger.warning(f"Error extracting text from PDF page {idx}: {e}")
                        extracted = None

                    if extracted and extracted.strip():
                        segments.append(
                            ExtractedSegment(
                                page_number=idx,
                                section_name=f"Page {idx}",
                                text=extracted.strip()
                            )
                        )
        except EmptyExtractionError:
            raise
        except Exception as e:
            raise FileCorruptedError(f"Failed to read PDF document. The file may be password-protected or corrupted: {str(e)}")

        if not segments:
            raise EmptyExtractionError(
                "PDF contains no extractable text. The file appears to be a scanned image or non-searchable document. "
                "Please upload a searchable text PDF or document with extractable text."
            )

    # 2. Modern Word Documents (.docx)
    elif file_type == "docx":
        try:
            doc = docx.Document(file_path)
            content_blocks: List[str] = []

            # Paragraphs
            for p in doc.paragraphs:
                p_text = p.text.strip()
                if p_text:
                    content_blocks.append(p_text)

            # Tables
            for table in doc.tables:
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_cells:
                        content_blocks.append(" | ".join(row_cells))

            if not content_blocks:
                raise EmptyExtractionError("DOCX document contains no readable text or tables.")

            full_docx_text = "\n\n".join(content_blocks)
            segments.append(
                ExtractedSegment(
                    page_number=1,
                    section_name="Document Body",
                    text=full_docx_text
                )
            )
        except EmptyExtractionError:
            raise
        except Exception as e:
            raise FileCorruptedError(f"Failed to read DOCX document. File may be corrupted or invalid: {str(e)}")

    # 3. Plain Text Documents (.txt)
    elif file_type == "txt":
        try:
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                text_content = f.read().strip()
        except Exception as e:
            raise FileCorruptedError(f"Failed to read text file: {str(e)}")

        if not text_content:
            raise EmptyExtractionError("Text file is empty.")

        segments.append(
            ExtractedSegment(
                page_number=1,
                section_name="Text Content",
                text=text_content
            )
        )

    # 4. JSON Policy & Audit Exports (.json)
    elif file_type == "json":
        try:
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                raw_json = json.load(f)

            formatted_text = json.dumps(raw_json, indent=2)
            if not formatted_text.strip() or formatted_text.strip() in ["{}", "[]", "null"]:
                raise EmptyExtractionError("JSON artifact contains no data.")

            segments.append(
                ExtractedSegment(
                    page_number=1,
                    section_name="JSON Configuration",
                    text=formatted_text
                )
            )
        except json.JSONDecodeError as e:
            raise FileCorruptedError(f"Invalid JSON syntax in compliance artifact: {str(e)}")

    # 5. CSV Audit Data (.csv)
    elif file_type == "csv":
        try:
            lines: List[str] = []
            with open(file_path, "r", encoding="utf-8", errors="replace", newline="") as f:
                reader = csv.reader(f)
                for row in reader:
                    if any(cell.strip() for cell in row):
                        lines.append(" | ".join(cell.strip() for cell in row))

            if not lines:
                raise EmptyExtractionError("CSV artifact contains no rows or data.")

            segments.append(
                ExtractedSegment(
                    page_number=1,
                    section_name="CSV Table",
                    text="\n".join(lines)
                )
            )
        except Exception as e:
            raise FileCorruptedError(f"Failed to parse CSV artifact: {str(e)}")

    else:
        raise UnsupportedFormatError(
            f"Unsupported file format '.{file_type}'. Supported formats: PDF, DOCX, TXT, JSON, CSV."
        )

    return segments


def extract_text(file_path: str, file_type: str) -> str:
    """
    Backward-compatible convenience function.
    Extracts text from file and returns single concatenated string.
    Raises ExtractionError if extraction produces no usable text.
    """
    segments = extract_document(file_path, file_type)
    full_text = "\n\n".join(seg.text for seg in segments if seg.text.strip()).strip()
    if not full_text:
        raise EmptyExtractionError("Extraction produced zero usable text.")
    return full_text
