import os
import re
import csv
import json
import logging
from pathlib import Path
from typing import List, Dict, Any, Optional
from sqlalchemy.orm import Session
from fastapi import HTTPException

import PyPDF2
import docx

from app.models.framework import Framework
from app.models.framework_version import FrameworkVersion
from app.models.control import Control
from app.models.control_relationship import ControlRelationship
from app.models.organization import Organization
from app.models.organization_framework import OrganizationFramework
from app.ai.embeddings import generate_embeddings
from app.ai.llm import get_llm

logger = logging.getLogger(__name__)

UPLOAD_DIR = Path("uploads/frameworks")
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)


def extract_raw_framework_text(file_path: str, file_type: str) -> str:
    """
    Extracts raw text from PDF, DOCX, TXT files with detailed error handling.
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"Framework file not found: {file_path}")

    text = ""
    file_type = file_type.lower().strip().lstrip(".")

    if file_type == "pdf":
        try:
            with open(file_path, "rb") as f:
                reader = PyPDF2.PdfReader(f)
                total_pages = len(reader.pages)
                if total_pages == 0:
                    raise ValueError("PDF file has 0 pages.")
                
                for idx, page in enumerate(reader.pages):
                    page_text = page.extract_text()
                    if page_text:
                        text += f"\n--- Page {idx + 1} ---\n" + page_text + "\n"
        except Exception as e:
            raise ValueError(f"Error parsing PDF file: {str(e)}")

    elif file_type in ["doc", "docx"]:
        try:
            doc = docx.Document(file_path)
            # Paragraphs
            for p in doc.paragraphs:
                if p.text.strip():
                    text += p.text + "\n"
            # Tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                    if row_text:
                        text += row_text + "\n"
        except Exception as e:
            raise ValueError(f"Error parsing DOCX file: {str(e)}")

    elif file_type == "txt":
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                text = f.read()
        except Exception as e:
            raise ValueError(f"Error reading TXT file: {str(e)}")

    else:
        raise ValueError(f"Unsupported document format for text extraction: {file_type}")

    return text.strip()


def normalize_rel_type(rel_type: str) -> str:
    """Normalizes relationship types to standard supported values."""
    rt = rel_type.lower().strip().replace("-", "_").replace(" ", "_")
    if rt in ["related_to", "related", "see_also", "reference", "references", "cross_reference", "cross_references"]:
        return "related_to"
    if rt in ["depends_on", "dependency", "dependencies", "requires", "prerequisite"]:
        return "depends_on"
    if rt in ["supports", "supported_by", "implements", "enables"]:
        return "supports"
    if rt in ["maps_to", "mapped_to", "mapping", "equivalent"]:
        return "maps_to"
    if rt in ["parent_of", "parent", "sub_controls", "children"]:
        return "parent_of"
    if rt in ["child_of", "child", "sub_control", "part_of"]:
        return "child_of"
    return "related_to"


def extract_explicit_cross_references(
    source_code: str,
    text: str,
    all_known_codes: set
) -> List[Dict[str, Any]]:
    """
    Scans a control's text for explicit cross-reference markers like:
    - 'Related controls: A.5.2, A.5.3'
    - 'See also: CC1.2'
    - 'Depends on: AC-2'
    - 'References: ...'
    - 'Parent Control: ...'
    Returns a list of raw relationship dicts.
    Only matches explicit keywords and validates target codes exist in all_known_codes.
    Never hallucinates or invents relationships.
    """
    if not text or len(text.strip()) == 0:
        return []

    relationships = []
    rel_patterns = [
        (r"(?i)(?:related\s+(?:controls?|requirements?)|related\s+to)[:\s]+([^\n\r]+)", "related_to"),
        (r"(?i)(?:depends\s+on|dependencies|requires|prerequisite)[:\s]+([^\n\r]+)", "depends_on"),
        (r"(?i)(?:supports|supported\s+by)[:\s]+([^\n\r]+)", "supports"),
        (r"(?i)(?:see\s+also|cross-?references?|references?)[:\s]+([^\n\r]+)", "related_to"),
        (r"(?i)(?:parent\s+control)[:\s]+([^\n\r]+)", "child_of"),
        (r"(?i)(?:sub-?controls?|child\s+controls?)[:\s]+([^\n\r]+)", "parent_of"),
        (r"(?i)(?:maps\s+to|mapped\s+to)[:\s]+([^\n\r]+)", "maps_to"),
    ]

    for pat, default_type in rel_patterns:
        for match in re.finditer(pat, text):
            line = match.group(1).strip()
            raw_tokens = re.split(r"[,;|\s]+(?:and\s+)?", line)
            for token in raw_tokens:
                cleaned = token.strip().rstrip(".,;:)").lstrip("(")
                if not cleaned:
                    continue
                for known in all_known_codes:
                    if cleaned.lower() == known.lower() and known.lower() != source_code.lower():
                        relationships.append({
                            "source_code": source_code,
                            "target_code": known,
                            "relationship_type": default_type,
                            "source_reference": f"Explicit reference in control text: '{match.group(0)[:80]}'",
                            "confidence": 1.0
                        })
    return relationships


def parse_structured_framework_file(file_path: str, file_type: str) -> Optional[List[Dict[str, Any]]]:
    """
    Parses structured JSON or CSV framework files directly into control records
    and extracts explicit control-to-control relationships if provided.
    """
    file_type = file_type.lower().strip().lstrip(".")
    controls: List[Dict[str, Any]] = []

    if file_type == "json":
        try:
            with open(file_path, "r", encoding="utf-8") as f:
                data = json.load(f)

            raw_controls = []
            if isinstance(data, list):
                raw_controls = data
            elif isinstance(data, dict):
                raw_controls = data.get("controls") or data.get("requirements") or data.get("items") or []

            for item in raw_controls:
                if isinstance(item, dict):
                    code = item.get("control_code") or item.get("code") or item.get("id") or item.get("control_id") or ""
                    title = item.get("title") or item.get("name") or item.get("control_title") or code
                    desc = item.get("description") or item.get("desc") or ""
                    req = item.get("requirement") or item.get("statement") or desc or title
                    cat = item.get("category") or item.get("domain") or item.get("family") or item.get("section") or "General"
                    guidance = item.get("guidance") or item.get("implementation_guidance") or ""
                    src_ref = item.get("source_reference") or item.get("ref") or str(code)

                    # Extract structured relationships if explicitly present
                    raw_rels = []
                    if "relationships" in item and isinstance(item["relationships"], list):
                        for r in item["relationships"]:
                            if isinstance(r, dict):
                                tgt = r.get("target_control_code") or r.get("target_code") or r.get("target") or r.get("target_control_id") or ""
                                rtype = r.get("relationship_type") or r.get("type") or "related_to"
                                ref = r.get("source_reference") or r.get("evidence") or r.get("reference") or "JSON document relationship"
                                conf = float(r.get("confidence") or r.get("mapping_confidence") or 1.0)
                                if tgt:
                                    raw_rels.append({
                                        "source_code": str(code).strip(),
                                        "target_code": str(tgt).strip(),
                                        "relationship_type": normalize_rel_type(str(rtype)),
                                        "source_reference": str(ref),
                                        "confidence": conf
                                    })

                    for field_key, rel_type in [
                        ("related_controls", "related_to"),
                        ("related_to", "related_to"),
                        ("depends_on", "depends_on"),
                        ("dependencies", "depends_on"),
                        ("parent_control", "child_of"),
                        ("sub_controls", "parent_of"),
                        ("supports", "supports"),
                        ("see_also", "related_to"),
                        ("references", "related_to"),
                        ("maps_to", "maps_to"),
                    ]:
                        val = item.get(field_key)
                        if val:
                            targets = val if isinstance(val, list) else re.split(r"[,;|\s]+", str(val))
                            for t in targets:
                                clean_t = str(t).strip().rstrip(".,;:)")
                                if clean_t:
                                    raw_rels.append({
                                        "source_code": str(code).strip(),
                                        "target_code": clean_t,
                                        "relationship_type": rel_type,
                                        "source_reference": f"Structured JSON attribute '{field_key}'",
                                        "confidence": 1.0
                                    })

                    if code and title:
                        controls.append({
                            "control_code": str(code).strip(),
                            "title": str(title).strip(),
                            "description": str(desc).strip(),
                            "requirement": str(req).strip(),
                            "category": str(cat).strip(),
                            "guidance": str(guidance).strip(),
                            "source_reference": str(src_ref).strip(),
                            "raw_relationships": raw_rels
                        })
            return controls if len(controls) > 0 else None
        except Exception as e:
            logger.warning(f"Failed to parse JSON structured controls: {e}")
            return None

    elif file_type == "csv":
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                reader = csv.DictReader(f)
                for row in reader:
                    low_row = {k.lower().strip().replace(" ", "_"): v.strip() for k, v in row.items() if k and v}
                    code = low_row.get("control_code") or low_row.get("code") or low_row.get("control_id") or low_row.get("id") or ""
                    title = low_row.get("title") or low_row.get("name") or low_row.get("control_name") or code
                    desc = low_row.get("description") or low_row.get("desc") or ""
                    req = low_row.get("requirement") or low_row.get("statement") or desc or title
                    cat = low_row.get("category") or low_row.get("domain") or low_row.get("family") or "General"
                    guidance = low_row.get("guidance") or low_row.get("implementation_guidance") or ""
                    src_ref = low_row.get("source_reference") or low_row.get("reference") or str(code)

                    raw_rels = []
                    for field_key, rel_type in [
                        ("related_controls", "related_to"),
                        ("related_to", "related_to"),
                        ("depends_on", "depends_on"),
                        ("dependencies", "depends_on"),
                        ("parent_control", "child_of"),
                        ("parent", "child_of"),
                        ("sub_controls", "parent_of"),
                        ("children", "parent_of"),
                        ("supports", "supports"),
                        ("see_also", "related_to"),
                        ("references", "related_to"),
                        ("mapped_controls", "maps_to"),
                        ("maps_to", "maps_to"),
                    ]:
                        val = low_row.get(field_key)
                        if val:
                            targets = re.split(r"[,;|\s]+", str(val))
                            for t in targets:
                                clean_t = str(t).strip().rstrip(".,;:)")
                                if clean_t:
                                    raw_rels.append({
                                        "source_code": str(code).strip(),
                                        "target_code": clean_t,
                                        "relationship_type": rel_type,
                                        "source_reference": f"Structured CSV column '{field_key}'",
                                        "confidence": 1.0
                                    })

                    if code and (title or req):
                        controls.append({
                            "control_code": str(code).strip(),
                            "title": str(title or code).strip(),
                            "description": str(desc).strip(),
                            "requirement": str(req).strip(),
                            "category": str(cat).strip(),
                            "guidance": str(guidance).strip(),
                            "source_reference": str(src_ref).strip(),
                            "raw_relationships": raw_rels
                        })
            return controls if len(controls) > 0 else None
        except Exception as e:
            logger.warning(f"Failed to parse CSV structured controls: {e}")
            return None

    return None


def parse_controls_from_text(raw_text: str, framework_name: str) -> List[Dict[str, Any]]:
    """
    Extracts real compliance controls and requirements from unstructured document text.
    Uses regex section detectors and falls back to LLM-assisted structural extraction.
    Strictly forbids hallucinating or generating fake controls.
    """
    if not raw_text or len(raw_text.strip()) < 20:
        return []

    controls: List[Dict[str, Any]] = []

    patterns = [
        # Match ISO Annex A: A.5.1, A.12.1.2
        r"(?:^|\n)(A\.\d+(?:\.\d+)*)\s+([^\n\r]+)(?:\n([\s\S]*?)(?=(?:\n[A-Z]\.\d+|\n--- Page|\Z)))",
        # Match NIST style: AC-1, PR.AC-1, IA-2(1), NIST CSF functions
        r"(?:^|\n)((?:[A-Z]{2,4}-[0-9]+(?:\([0-9]+\))?)|(?:[A-Z]{2}\.[A-Z]{2}-[0-9]+))\s+([^\n\r]+)(?:\n([\s\S]*?)(?=(?:\n(?:[A-Z]{2,4}-[0-9]+|[A-Z]{2}\.[A-Z]{2}-[0-9]+)|\n--- Page|\Z)))",
        # Match SOC 2 style: CC1.1, CC6.2, A1.1
        r"(?:^|\n)(CC\d+\.\d+|[A-Z]\d+\.\d+)\s+([^\n\r]+)(?:\n([\s\S]*?)(?=(?:\n(?:CC\d+\.\d+|[A-Z]\d+\.\d+)|\n--- Page|\Z)))",
        # Match Generic "Control X.Y" or "Section X.Y"
        r"(?:^|\n)(?:Control|Section|Req|Requirement)\s*([0-9]+(?:\.[0-9]+)+)[:\s]+([^\n\r]+)(?:\n([\s\S]*?)(?=(?:\n(?:Control|Section|Req|Requirement)\s*[0-9]+|\n--- Page|\Z)))",
        # Match Numbered sections: "1.1 Access Control" / "5.1.2 Security Policy"
        r"(?:^|\n)([0-9]+(?:\.[0-9]+){1,3})\s+([A-Z][^\n\r]{3,80})(?:\n([\s\S]*?)(?=(?:\n[0-9]+(?:\.[0-9]+){1,3}\s+[A-Z]|\n--- Page|\Z)))"
    ]

    for pat in patterns:
        matches = list(re.finditer(pat, raw_text, re.MULTILINE))
        if len(matches) >= 3:
            for m in matches:
                code = m.group(1).strip()
                title = m.group(2).strip()
                body = m.group(3).strip() if len(m.groups()) >= 3 and m.group(3) else ""

                category = "General"
                if "." in code:
                    prefix = code.split(".")[0]
                    category = f"Domain {prefix}"
                elif "-" in code:
                    prefix = code.split("-")[0]
                    category = f"Family {prefix}"

                requirement = body if body else title

                controls.append({
                    "control_code": code,
                    "title": title,
                    "description": body[:500] if body else title,
                    "requirement": requirement,
                    "category": category,
                    "guidance": "",
                    "source_reference": f"{framework_name} {code}",
                    "raw_relationships": []
                })
            break

    if len(controls) >= 3:
        return controls

    try:
        llm = get_llm()
        if llm:
            chunks = []
            paragraphs = raw_text.split("\n\n")
            current_chunk = ""
            for p in paragraphs:
                if len(current_chunk) + len(p) < 3000:
                    current_chunk += p + "\n\n"
                else:
                    if current_chunk.strip():
                        chunks.append(current_chunk.strip())
                    current_chunk = p + "\n\n"
            if current_chunk.strip():
                chunks.append(current_chunk.strip())

            extracted_from_llm = []
            for chunk in chunks[:8]:
                prompt = (
                    f"You are a compliance parser. Extract all compliance controls and requirements "
                    f"from the following text verbatim.\n"
                    f"CRITICAL RULES:\n"
                    f"1. ONLY extract controls that are EXPLICITLY present in the text below.\n"
                    f"2. DO NOT invent, hallucinate, or summarize missing controls.\n"
                    f"3. If a control explicitly states relationships to other controls (e.g. 'Related controls:', 'Depends on:', 'See also:'), include them in a \"relationships\" array: [{{\"target_control_code\": \"...\", \"relationship_type\": \"related_to|depends_on|supports|parent_of|child_of\", \"evidence\": \"...\"}}]. NEVER guess or invent relationships.\n"
                    f"4. If no controls are found in this snippet, respond with an empty JSON array: []\n"
                    f"5. Output MUST be valid JSON only, formatted as a JSON array of objects with keys: "
                    f"\"control_code\", \"title\", \"requirement\", \"category\", \"guidance\", \"relationships\".\n\n"
                    f"TEXT SNIPPET:\n{chunk}\n\n"
                    f"JSON RESPONSE:"
                )
                
                res = llm.invoke(prompt)
                content = res.content.strip()
                if content.startswith("```json"):
                    content = content[7:]
                if content.startswith("```"):
                    content = content[3:]
                if content.endswith("```"):
                    content = content[:-3]
                content = content.strip()

                try:
                    parsed = json.loads(content)
                    if isinstance(parsed, list):
                        for c in parsed:
                            if isinstance(c, dict) and c.get("control_code") and c.get("title"):
                                raw_rels = []
                                if "relationships" in c and isinstance(c["relationships"], list):
                                    for r in c["relationships"]:
                                        if isinstance(r, dict) and r.get("target_control_code"):
                                            raw_rels.append({
                                                "source_code": str(c["control_code"]).strip(),
                                                "target_code": str(r["target_control_code"]).strip(),
                                                "relationship_type": normalize_rel_type(str(r.get("relationship_type", "related_to"))),
                                                "source_reference": str(r.get("evidence") or "LLM extracted cross-reference"),
                                                "confidence": 0.95
                                            })
                                extracted_from_llm.append({
                                    "control_code": str(c["control_code"]).strip(),
                                    "title": str(c["title"]).strip(),
                                    "description": str(c.get("requirement") or c["title"]).strip(),
                                    "requirement": str(c.get("requirement") or c["title"]).strip(),
                                    "category": str(c.get("category") or "General").strip(),
                                    "guidance": str(c.get("guidance") or "").strip(),
                                    "source_reference": f"{framework_name} {c['control_code']}",
                                    "raw_relationships": raw_rels
                                })
                except Exception:
                    continue

            if len(extracted_from_llm) > 0:
                unique_controls = {}
                for c in extracted_from_llm:
                    unique_controls[c["control_code"]] = c
                return list(unique_controls.values())

    except Exception as e:
        logger.warning(f"LLM control extraction failed: {e}")

    return controls


def process_framework_document(
    db: Session,
    framework_id: int,
    framework_version_id: int
) -> Dict[str, Any]:
    """
    Main pipeline for processing an uploaded compliance framework:
    1. Reads and extracts text from the real stored file.
    2. Identifies and parses real controls & requirements.
    3. Stores controls in PostgreSQL.
    4. Automatically extracts, validates, and stores explicit control-to-control relationships.
    5. Generates vector embeddings for each control using nomic-embed-text / Ollama.
    6. Updates framework status and enrolls active organizations.
    """
    framework = db.query(Framework).filter(Framework.id == framework_id).first()
    if not framework:
        raise HTTPException(status_code=404, detail="Framework not found")

    version = db.query(FrameworkVersion).filter(FrameworkVersion.id == framework_version_id).first()
    if not version:
        raise HTTPException(status_code=404, detail="Framework version not found")

    if not framework.file_path or not os.path.exists(framework.file_path):
        framework.status = "failed"
        framework.error_message = f"Framework file not found at: {framework.file_path}"
        db.commit()
        raise HTTPException(status_code=400, detail="Framework source file does not exist on disk")

    framework.status = "processing"
    db.commit()

    try:
        ext = framework.file_name.split(".")[-1].lower() if framework.file_name and "." in framework.file_name else "txt"
        
        # Step 1: Check for structured formats (JSON/CSV)
        extracted_controls = parse_structured_framework_file(framework.file_path, ext)

        # Step 2: If not structured, extract raw text from PDF/DOCX/TXT
        if not extracted_controls:
            raw_text = extract_raw_framework_text(framework.file_path, ext)
            
            if not raw_text or len(raw_text.strip()) == 0:
                framework.status = "review_required"
                framework.error_message = (
                    "Document text extraction yielded 0 characters. "
                    "The document may be scanned or image-only (OCR is required)."
                )
                db.commit()
                return {
                    "framework_id": framework.id,
                    "version_id": version.id,
                    "name": framework.name,
                    "code": framework.code,
                    "status": "review_required",
                    "total_controls": 0,
                    "relationships_detected": 0,
                    "relationships_created": 0,
                    "relationships_rejected": 0,
                    "message": framework.error_message
                }

            # Step 3: Parse controls from extracted text
            extracted_controls = parse_controls_from_text(raw_text, framework.name)

        if not extracted_controls or len(extracted_controls) == 0:
            framework.status = "failed"
            framework.error_message = (
                "Could not identify structured controls in the uploaded document. "
                "Ensure the document contains recognizable compliance control sections."
            )
            db.commit()
            return {
                "framework_id": framework.id,
                "version_id": version.id,
                "name": framework.name,
                "code": framework.code,
                "status": "failed",
                "total_controls": 0,
                "relationships_detected": 0,
                "relationships_created": 0,
                "relationships_rejected": 0,
                "message": framework.error_message
            }

        # Step 4: Clean prior controls & control relationships for this framework_version_id for idempotency
        existing_ctrl_ids = [c.id for c in db.query(Control.id).filter(Control.framework_version_id == version.id).all()]
        if existing_ctrl_ids:
            db.query(ControlRelationship).filter(
                (ControlRelationship.source_control_id.in_(existing_ctrl_ids)) |
                (ControlRelationship.target_control_id.in_(existing_ctrl_ids))
            ).delete(synchronize_session=False)
        db.query(Control).filter(Control.framework_version_id == version.id).delete()
        db.flush()

        # Step 5: Generate vector embeddings for controls in batches
        texts_to_embed = [
            f"Control {c['control_code']}: {c['title']}\n"
            f"Category: {c.get('category', 'General')}\n"
            f"Requirement: {c.get('requirement', '')}\n"
            f"Guidance: {c.get('guidance', '')}".strip()
            for c in extracted_controls
        ]

        logger.info(f"Generating embeddings for {len(texts_to_embed)} controls of {framework.name}...")
        try:
            embeddings = generate_embeddings(texts_to_embed)
        except Exception as emb_err:
            logger.warning(f"Embedding generation via Ollama failed: {emb_err}. Storing controls without embeddings.")
            embeddings = [None] * len(extracted_controls)

        # Step 6: Store Controls in PostgreSQL
        stored_controls = []
        for idx, ctrl_data in enumerate(extracted_controls):
            emb = embeddings[idx] if idx < len(embeddings) else None
            control_row = Control(
                framework_version_id=version.id,
                control_code=ctrl_data["control_code"],
                title=ctrl_data["title"],
                description=ctrl_data.get("description") or ctrl_data["title"],
                requirement=ctrl_data.get("requirement") or ctrl_data.get("description") or ctrl_data["title"],
                category=ctrl_data.get("category") or "General",
                guidance=ctrl_data.get("guidance"),
                source_reference=ctrl_data.get("source_reference") or f"{framework.code} {ctrl_data['control_code']}",
                is_active=True,
                embedding=emb
            )
            db.add(control_row)
            stored_controls.append(control_row)

        db.flush()  # Flush so that stored_controls get their generated database IDs

        # Step 7: Automatic Control-to-Control Relationship Extraction & Validation
        code_to_control_id = {c.control_code.upper().strip(): c.id for c in stored_controls}
        all_known_codes = set(code_to_control_id.keys())

        all_candidate_rels: List[Dict[str, Any]] = []

        for c_data in extracted_controls:
            src_code = str(c_data.get("control_code", "")).strip()

            # 1. Structured relationships from JSON/CSV/LLM
            if "raw_relationships" in c_data and isinstance(c_data["raw_relationships"], list):
                for r in c_data["raw_relationships"]:
                    all_candidate_rels.append(r)

            # 2. Text scanning for explicit cross-references in requirement, description, and guidance
            full_text = f"{c_data.get('title', '')} {c_data.get('requirement', '')} {c_data.get('description', '')} {c_data.get('guidance', '')}"
            text_rels = extract_explicit_cross_references(src_code, full_text, all_known_codes)
            all_candidate_rels.extend(text_rels)

        # Validate relationships and insert into control_relationships
        detected_count = len(all_candidate_rels)
        created_count = 0
        rejected_count = 0
        seen_pairs = set()

        for rel in all_candidate_rels:
            src_code = str(rel.get("source_code", "")).strip().upper()
            tgt_code = str(rel.get("target_code", "")).strip().upper()
            rel_type = normalize_rel_type(str(rel.get("relationship_type", "related_to")))
            src_ref = rel.get("source_reference")
            conf = float(rel.get("confidence", 1.0)) if rel.get("confidence") is not None else 1.0

            src_id = code_to_control_id.get(src_code)
            tgt_id = code_to_control_id.get(tgt_code)

            # Validate that both controls belong to this framework and exist
            if not src_id or not tgt_id:
                rejected_count += 1
                continue

            # Never self-link
            if src_id == tgt_id:
                rejected_count += 1
                continue

            # Prevent duplicates
            pair_key = (src_id, tgt_id, rel_type)
            if pair_key in seen_pairs:
                rejected_count += 1
                continue

            seen_pairs.add(pair_key)

            rel_record = ControlRelationship(
                source_control_id=src_id,
                target_control_id=tgt_id,
                relationship_type=rel_type,
                source_reference=str(src_ref)[:255] if src_ref else f"{framework.code} cross-reference",
                mapping_confidence=conf
            )
            db.add(rel_record)
            created_count += 1

        db.flush()

        # Step 8: Update counts and status
        total_count = len(stored_controls)
        framework.total_controls = total_count
        framework.status = "completed"
        framework.error_message = None
        version.total_controls = total_count
        version.status = "active"

        # Step 9: Auto-associate framework with existing tenant organizations
        all_orgs = db.query(Organization).all()
        for org in all_orgs:
            existing_ofw = db.query(OrganizationFramework).filter(
                OrganizationFramework.organization_id == org.id,
                OrganizationFramework.framework_id == framework.id
            ).first()
            if not existing_ofw:
                ofw = OrganizationFramework(
                    organization_id=org.id,
                    framework_id=framework.id,
                    framework_version_id=version.id,
                    status="active"
                )
                db.add(ofw)

        db.commit()
        db.refresh(framework)
        db.refresh(version)

        logger.info(
            f"Successfully ingested framework '{framework.name}' with {total_count} controls "
            f"and {created_count} control relationships."
        )

        return {
            "framework_id": framework.id,
            "version_id": version.id,
            "name": framework.name,
            "code": framework.code,
            "status": "completed",
            "total_controls": total_count,
            "relationships_detected": detected_count,
            "relationships_created": created_count,
            "relationships_rejected": rejected_count,
            "message": (
                f"Successfully ingested {total_count} controls and {created_count} "
                f"control relationships with pgvector embeddings."
            )
        }

    except Exception as e:
        db.rollback()
        framework.status = "failed"
        framework.error_message = str(e)
        db.commit()
        logger.error(f"Framework ingestion pipeline failed: {str(e)}", exc_info=True)
        raise HTTPException(status_code=500, detail=f"Framework ingestion failed: {str(e)}")

